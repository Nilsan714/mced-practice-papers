"""Excel performance tracker (formulas depend on actual candidate responses entered by the tutor)."""
import os, sys, json
HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
sys.path.insert(0, os.path.join(HERE, 'content'))
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter as L
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import CellIsRule
from build_docs import load_papers, OUT

MAXQ = 30
HDR = PatternFill('solid', fgColor='1F3864')
INP = PatternFill('solid', fgColor='FFF2CC')
CALC = PatternFill('solid', fgColor='E2EFDA')
WH = Font(color='FFFFFF', bold=True)
THIN = Side(style='thin', color='BFBFBF')
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
BASE_SECTIONS = [('A Bookkeeping & Basic Accounting', 1, 12, 'Papers 2-5'), ('B Taxation', 13, 18, 'Paper 6'), ('C Budget & Financial Discipline', 19, 23, 'Paper 7 (+5 BRS)'),
                 ('D Banking', 24, 29, 'Paper 8'), ('E Computer & Modern Tools', 30, 35, 'Paper 9'), ('F Current Affairs / GA', 36, 40, 'Paper 10')]
BASE_KEY = 'ABABBABBCBBCBBCBBABAABBAABAAABBBBCAAABBC'  # matches the supplied Teacher Answer Key (Quick Scoring table)


def hdr(ws, row, vals, widths=None):
    for i, v in enumerate(vals, 1):
        c = ws.cell(row=row, column=i, value=v)
        c.fill, c.font, c.border = HDR, WH, BOX
        c.alignment = Alignment(wrap_text=True, vertical='center', horizontal='center')
    if widths:
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[L(i)].width = w


def build(path):
    papers = load_papers()
    wb = Workbook()
    # ---------- README
    ws = wb.active
    ws.title = 'README'
    lines = ['MCED District Account Assistant — Practice Performance Tracker',
             'Practice scoring rule (NOT an official MCED scheme): 1 mark per correct answer, 0 for wrong/unanswered, no negative marking (change only if you deliberately configure another rule).',
             '',
             'HOW TO USE (tutor only — this workbook contains the answer key):',
             '1. Candidate solves the question-only paper without the answer key; tutor records the answers.',
             '2. Sheet "Attempts": add one row per attempt. Fill the yellow cells: Attempt ID (e.g. A001), Candidate, Date, Paper No (1-11), then the candidate\'s responses A/B/C/D in Q1..Q30 (leave blank = unanswered).',
             '3. Green cells calculate automatically: attempted / correct / incorrect / unanswered / score / %.',
             '4. Sheet "Topic_Accuracy": topic-wise accuracy per attempt, weakest subtopic and recommended remedial lesson.',
             '5. Sheet "Error_Log": for each wrong answer, choose an error category (Conceptual / Calculation / Terminology / Careless).',
             '6. Teach the recommended lesson (Remediation Notes), give a short remedial quiz, then record the retest as a new Attempt and link both IDs in sheet "Retest".',
             '7. Sheet "Baseline_40": enter the candidate\'s 40 baseline answers for section-wise baseline analysis.',
             '',
             'Do NOT label the candidate weak/strong from total score alone; use topic-wise accuracy (suggested guide, not an official cutoff): <50% re-teach, 50-74% practice, >=75% relatively secure. Use topic % only when the subtopic has at least 3 questions.',
             'Rows 2..200 of Attempts are pre-formulated. Copy a row down if you need more.']
    for i, t in enumerate(lines, 1):
        ws.cell(row=i, column=1, value=t).alignment = Alignment(wrap_text=True, vertical='top')
    ws['A1'].font = Font(bold=True, size=14)
    ws.column_dimensions['A'].width = 130
    # ---------- Papers
    wp = wb.create_sheet('Papers')
    hdr(wp, 1, ['Paper No', 'Title', 'Questions', 'Minutes', 'Sub1', 'Sub2', 'Sub3', 'Sub4', 'Sub5', 'Sub6', 'Les1', 'Les2', 'Les3', 'Les4', 'Les5', 'Les6'],
        [9, 48, 10, 9] + [34] * 6 + [34] * 6)
    for r, p in enumerate(papers, 2):
        row = [p['no'], p['title'], len(p['questions']), p['duration']] + [(p['subtopics'] + [''] * 6)[i] for i in range(6)] + [(p['lessons'] + [''] * 6)[i] for i in range(6)]
        for c, v in enumerate(row, 1):
            wp.cell(row=r, column=c, value=v).border = BOX
    # ---------- Key (tutor only) and SubMap
    wk = wb.create_sheet('Key')
    hdr(wk, 1, ['Paper', 'Q', 'Answer', 'Subtopic#', 'Subtopic', 'Level', 'ID(P*100+Q)'], [8, 6, 8, 10, 40, 14, 12])
    wm = wb.create_sheet('SubMap')
    hdr(wm, 1, ['Paper'] + ['Q%d' % i for i in range(1, MAXQ + 1)], [8] + [5] * MAXQ)
    r = 2
    for pi, p in enumerate(papers):
        wm.cell(row=pi + 2, column=1, value=p['no'])
        for qi, q in enumerate(p['questions'], 1):
            wk.append([p['no'], qi, q['ans'], q['sub'] + 1, p['subtopics'][q['sub']], {'F': 'Foundational', 'I': 'Intermediate', 'A': 'Application'}[q['lvl']], p['no'] * 100 + qi])
            wm.cell(row=pi + 2, column=qi + 1, value=q['sub'] + 1)
    nkey = wk.max_row
    # ---------- Attempts
    wa = wb.create_sheet('Attempts')
    heads = ['Attempt ID', 'Candidate', 'Date', 'Paper No', 'Topic', 'Total Qs', 'Attempted', 'Correct', 'Incorrect', 'Unanswered', 'Score', '%'] + ['Q%d' % i for i in range(1, MAXQ + 1)]
    hdr(wa, 1, heads, [11, 22, 12, 9, 40, 8, 10, 8, 9, 11, 7, 8] + [5] * MAXQ)
    flag0 = 12 + MAXQ + 2  # flag columns start (after a spacer column)
    wa.cell(row=1, column=flag0 - 1, value='→ flags')
    for i in range(1, MAXQ + 1):
        c = wa.cell(row=1, column=flag0 + i - 1, value='f%d' % i)
        c.fill, c.font = HDR, WH
        wa.column_dimensions[L(flag0 + i - 1)].width = 4
    NR = 200
    dvp = DataValidation(type='whole', operator='between', formula1='1', formula2='11', allow_blank=True)
    dvr = DataValidation(type='list', formula1='"A,B,C,D"', allow_blank=True)
    wa.add_data_validation(dvp)
    wa.add_data_validation(dvr)
    dvp.add('D2:D%d' % NR)
    dvr.add('M2:%s%d' % (L(12 + MAXQ), NR))
    for r in range(2, NR + 1):
        for c in range(1, 5):
            wa.cell(row=r, column=c).fill = INP
        for c in range(13, 13 + MAXQ):
            wa.cell(row=r, column=c).fill = INP
        wa['E%d' % r] = '=IF($D%d="","",IFERROR(INDEX(Papers!$B$2:$B$12,MATCH($D%d,Papers!$A$2:$A$12,0)),"?"))' % (r, r)
        wa['F%d' % r] = '=IF($D%d="","",IFERROR(INDEX(Papers!$C$2:$C$12,MATCH($D%d,Papers!$A$2:$A$12,0)),0))' % (r, r)
        rng = 'M%d:%s%d' % (r, L(12 + MAXQ), r)
        fr = '%s%d:%s%d' % (L(flag0), r, L(flag0 + MAXQ - 1), r)
        wa['G%d' % r] = '=IF($D%d="","",COUNTIF(%s,"?*"))' % (r, rng)
        wa['H%d' % r] = '=IF($D%d="","",SUM(%s))' % (r, fr)
        wa['I%d' % r] = '=IF($D%d="","",G%d-H%d)' % (r, r, r)
        wa['J%d' % r] = '=IF($D%d="","",F%d-G%d)' % (r, r, r)
        wa['K%d' % r] = '=IF($D%d="","",H%d)' % (r, r)
        wa['L%d' % r] = '=IF($D%d="","",IF(F%d=0,"",ROUND(H%d/F%d*100,1)))' % (r, r, r, r)
        for c in range(5, 13):
            wa.cell(row=r, column=c).fill = CALC
        for q in range(1, MAXQ + 1):
            resp = '%s%d' % (L(12 + q), r)
            wa.cell(row=r, column=flag0 + q - 1, value=(
                '=IF(OR($D%d="",%s="",%d>$F%d),0,IF(UPPER(TRIM(%s))=IFERROR(INDEX(Key!$C$2:$C$%d,MATCH($D%d*100+%d,Key!$G$2:$G$%d,0)),"#"),1,0))'
                % (r, resp, q, r, resp, nkey, r, q, nkey)))
    wa.freeze_panes = 'E2'
    # ---------- Topic_Accuracy
    wt = wb.create_sheet('Topic_Accuracy')
    th = ['Attempt ID', 'Candidate', 'Paper']
    for s in range(1, 7):
        th += ['S%d name' % s, 'S%d correct' % s, 'S%d total' % s, 'S%d %%' % s]
    th += ['Weakest subtopic', 'Weakest %', 'Recommended remedial lesson', 'Guide (not official)']
    hdr(wt, 1, th, [11, 20, 7] + [26, 8, 8, 8] * 6 + [30, 10, 36, 22])
    for r in range(2, NR + 1):
        wt['A%d' % r] = '=IF(Attempts!A%d="","",Attempts!A%d)' % (r, r)
        wt['B%d' % r] = '=IF(Attempts!A%d="","",Attempts!B%d)' % (r, r)
        wt['C%d' % r] = '=IF(Attempts!A%d="","",Attempts!D%d)' % (r, r)
        fr = 'Attempts!$%s%d:$%s%d' % (L(flag0), r, L(flag0 + MAXQ - 1), r)
        pcts = []
        for s in range(1, 7):
            base = 4 + (s - 1) * 4
            cn, cc, ct, cp = L(base), L(base + 1), L(base + 2), L(base + 3)
            wt['%s%d' % (cn, r)] = '=IF($A%d="","",IFERROR(INDEX(Papers!$%s$2:$%s$12,MATCH($C%d,Papers!$A$2:$A$12,0)),""))' % (r, L(4 + s), L(4 + s), r)
            wt['%s%d' % (ct, r)] = '=IF($A%d="","",IFERROR(COUNTIF(INDEX(SubMap!$B$2:$AE$12,MATCH($C%d,SubMap!$A$2:$A$12,0),0),%d),0))' % (r, r, s)
            wt['%s%d' % (cc, r)] = '=IF($A%d="","",IFERROR(SUMPRODUCT((INDEX(SubMap!$B$2:$AE$12,MATCH($C%d,SubMap!$A$2:$A$12,0),0)=%d)*%s),0))' % (r, r, s, fr)
            wt['%s%d' % (cp, r)] = '=IF(OR($A%d="",%s%d=0),"",ROUND(%s%d/%s%d*100,1))' % (r, ct, r, cc, r, ct, r)
            pcts.append('%s%d' % (cp, r))
        prange = ','.join(pcts)
        wt['AB%d' % r] = ('=IF($A%d="","",IFERROR(INDEX(CHOOSE({1,2,3,4,5,6},%s),MATCH(MIN(%s),CHOOSE({1,2,3,4,5,6},%s),0)),""))'
                         % (r, ','.join('%s%d' % (L(4 + i * 4), r) for i in range(6)), prange, prange))
        wt['AC%d' % r] = '=IF($A%d="","",IF(COUNT(%s)=0,"",MIN(%s)))' % (r, prange, prange)
        wt['AD%d' % r] = ('=IF($A%d="","",IFERROR(INDEX(CHOOSE({1,2,3,4,5,6},%s),MATCH(MIN(%s),CHOOSE({1,2,3,4,5,6},%s),0)),""))'
                          % (r, ','.join('IFERROR(INDEX(Papers!$%s$2:$%s$12,MATCH($C%d,Papers!$A$2:$A$12,0)),"")' % (L(10 + i), L(10 + i), r) for i in range(1, 7)), prange, prange))
        wt['AE%d' % r] = '=IF(AC%d="","",IF(AC%d<50,"Re-teach",IF(AC%d<75,"Practice more","Relatively secure")))' % (r, r, r)
    wt.freeze_panes = 'D2'
    # ---------- Error_Log
    we = wb.create_sheet('Error_Log')
    hdr(we, 1, ['Attempt ID', 'Q No', 'Paper', 'Candidate answer', 'Correct answer', 'Subtopic', 'Error category', 'Recommended remedial lesson', 'Tutor note', 'Remedial lesson done? (Y/N)'],
        [11, 7, 7, 12, 12, 34, 22, 40, 36, 14])
    dve = DataValidation(type='list', formula1='"Conceptual misunderstanding,Calculation error,Terminology error,Careless mistake"', allow_blank=True)
    we.add_data_validation(dve)
    dve.add('G2:G500')
    for r in range(2, 501):
        we['A%d' % r].fill = we['B%d' % r].fill = we['G%d' % r].fill = we['I%d' % r].fill = we['J%d' % r].fill = INP
        we['C%d' % r] = '=IF(A%d="","",IFERROR(INDEX(Attempts!$D$2:$D$%d,MATCH(A%d,Attempts!$A$2:$A$%d,0)),"?"))' % (r, NR, r, NR)
        we['D%d' % r] = '=IF(OR(A%d="",B%d=""),"",IFERROR(INDEX(Attempts!$M$2:$%s$%d,MATCH(A%d,Attempts!$A$2:$A$%d,0),B%d),""))' % (r, r, L(12 + MAXQ), NR, r, NR, r)
        we['E%d' % r] = '=IF(OR(C%d="",B%d=""),"",IFERROR(INDEX(Key!$C$2:$C$%d,MATCH(C%d*100+B%d,Key!$G$2:$G$%d,0)),""))' % (r, r, nkey, r, r, nkey)
        we['F%d' % r] = '=IF(OR(C%d="",B%d=""),"",IFERROR(INDEX(Key!$E$2:$E$%d,MATCH(C%d*100+B%d,Key!$G$2:$G$%d,0)),""))' % (r, r, nkey, r, r, nkey)
        we['H%d' % r] = ('=IF(OR(C%d="",B%d=""),"",IFERROR(INDEX(Papers!$K$2:$P$12,MATCH(C%d,Papers!$A$2:$A$12,0),INDEX(Key!$D$2:$D$%d,MATCH(C%d*100+B%d,Key!$G$2:$G$%d,0))),""))'
                         % (r, r, r, nkey, r, r, nkey))
    # ---------- Retest
    wr = wb.create_sheet('Retest')
    hdr(wr, 1, ['Original Attempt ID', 'Retest Attempt ID', 'Candidate', 'Paper (orig)', 'Original %', 'Retest %', 'Improvement (percentage points)', 'Retest paper', 'Note'], [16, 16, 20, 11, 11, 11, 18, 11, 30])
    for r in range(2, 202):
        wr['A%d' % r].fill = wr['B%d' % r].fill = wr['I%d' % r].fill = INP
        wr['C%d' % r] = '=IF(A%d="","",IFERROR(INDEX(Attempts!$B$2:$B$%d,MATCH(A%d,Attempts!$A$2:$A$%d,0)),"?"))' % (r, NR, r, NR)
        wr['D%d' % r] = '=IF(A%d="","",IFERROR(INDEX(Attempts!$D$2:$D$%d,MATCH(A%d,Attempts!$A$2:$A$%d,0)),"?"))' % (r, NR, r, NR)
        wr['E%d' % r] = '=IF(A%d="","",IFERROR(INDEX(Attempts!$L$2:$L$%d,MATCH(A%d,Attempts!$A$2:$A$%d,0)),""))' % (r, NR, r, NR)
        wr['F%d' % r] = '=IF(B%d="","",IFERROR(INDEX(Attempts!$L$2:$L$%d,MATCH(B%d,Attempts!$A$2:$A$%d,0)),""))' % (r, NR, r, NR)
        wr['G%d' % r] = '=IF(OR(E%d="",F%d=""),"",F%d-E%d)' % (r, r, r, r)
        wr['H%d' % r] = '=IF(B%d="","",IFERROR(INDEX(Attempts!$D$2:$D$%d,MATCH(B%d,Attempts!$A$2:$A$%d,0)),"?"))' % (r, NR, r, NR)
    # ---------- Baseline_40
    wbz = wb.create_sheet('Baseline_40')
    wbz['A1'], wbz['B1'] = 'Candidate', None
    wbz['A2'], wbz['B2'] = 'Date', None
    wbz['B1'].fill = wbz['B2'].fill = INP
    hdr(wbz, 4, ['Q', 'Section #', 'Correct answer', 'Candidate answer', 'Correct (1/0)'], [10, 40, 14, 16, 14])
    sec_of = {}
    for si, (nm, a, b, ls) in enumerate(BASE_SECTIONS, 1):
        for q in range(a, b + 1):
            sec_of[q] = si
    for q in range(1, 41):
        r = 4 + q
        wbz['A%d' % r] = q
        wbz['B%d' % r] = sec_of[q]
        wbz['C%d' % r] = BASE_KEY[q - 1]
        wbz['D%d' % r].fill = INP
        wbz['E%d' % r] = '=IF(D%d="",0,IF(UPPER(TRIM(D%d))=C%d,1,0))' % (r, r, r)
    dvb = DataValidation(type='list', formula1='"A,B,C,D"', allow_blank=True)
    wbz.add_data_validation(dvb)
    dvb.add('D5:D44')
    for i, h in enumerate(['Section', 'Questions', 'Correct', '%', 'Guide (not official)', 'Next papers'], 7):
        c = wbz.cell(row=4, column=i, value=h)
        c.fill, c.font = HDR, WH
    for i, w in zip(range(7, 13), [34, 8, 9, 8, 22, 18]):
        wbz.column_dimensions[L(i)].width = w
    for si, (nm, a, b, ls) in enumerate(BASE_SECTIONS, 1):
        r = 4 + si
        wbz['G%d' % r] = nm
        wbz['H%d' % r] = b - a + 1
        wbz['I%d' % r] = '=SUMIF($B$5:$B$44,%d,$E$5:$E$44)' % si
        wbz['J%d' % r] = '=ROUND(I%d/H%d*100,1)' % (r, r)
        wbz['K%d' % r] = '=IF(COUNTA($D$5:$D$44)=0,"",IF(J%d<50,"Needs re-teaching",IF(J%d<75,"Developing","Relatively strong")))' % (r, r)
        wbz['L%d' % r] = ls
    wbz['G12'], wbz['H12'] = 'TOTAL', '=SUM(H5:H10)'
    wbz['I12'], wbz['J12'] = '=SUM(I5:I10)', '=ROUND(I12/H12*100,1)'
    wbz['G14'] = 'Baseline key = the supplied Teacher Answer Key (Quick Scoring). Diagnostic only; no official cutoff. Decide teaching from section/question evidence, not the total.'
    wbz['G14'].alignment = Alignment(wrap_text=True)
    wbz.column_dimensions['G'].width = 44
    wb.save(path)
    return path


if __name__ == '__main__':
    out = os.path.join(OUT, 'MCED_Practice_Performance_Tracker.xlsx')
    print(build(out))
